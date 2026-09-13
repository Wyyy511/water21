from __future__ import annotations
import io, re, json, difflib
from pathlib import Path
from typing import Dict, List, Tuple
import pandas as pd
from docx import Document
from pypdf import PdfReader

from core.paths import ENGINE_DIR
import importlib.util

spec=importlib.util.spec_from_file_location('baseline_v2', ENGINE_DIR/'baseline_water_risk_v2.py')
b=importlib.util.module_from_spec(spec); spec.loader.exec_module(b)

CANONICAL=[
 'enterprise','material','node_id','node_name','purchase_weight','year',
 'purchase_quantity','purchase_unit','procurement_basis',
 'WS','WS_unit','SV','SV_unit','DR','BWD','DYS','DYS_unit','CTS','OA',
 'theta_WS','theta_DR','theta_SV',
 'source_type','confidence','data_origin','is_simulation','human_confirmed',
 'unknown_flag','evidence'
]
ALIASES={
 'enterprise':['enterprise','enterprise_name','enterprise_id','企业','企业名称','公司','company'],
 'material':['material','material_id','原材料','材料','原材料标识','raw material','commodity'],
 'node_id':['node_id','supplier_node_id','供应节点编码','节点编码','节点id','id'],
 'node_name':['node_name','供应节点','节点','产区','地区','location','supplier location','source country','country'],
 'purchase_weight':['purchase_weight','w','采购权重','采购占比','采购比例','purchase %','purchase_share','share','weight'],
 'purchase_quantity':['purchase_quantity','采购量','采购数量','quantity','volume'],
 'purchase_unit':['purchase_unit','采购单位','数量单位','unit'],
 'procurement_basis':['procurement_basis','采购口径','basis'],
 'year':['year','procurement_year','采购年份','年份','参考期','period'],
 'WS':['WS','ws','water_stress','水压力'],
 'WS_unit':['WS_unit','ws_unit','WS尺度','水压力尺度'],
 'SV':['SV','sv','seasonal_variability','季节波动'],
 'SV_unit':['SV_unit','sv_unit','SV尺度','季节波动尺度'],
 'DR':['DR','dr','drought','历史干旱'],
 'BWD':['BWD','bwd','蓝水依赖'],
 'DYS':['DYS','dys','产出敏感性'],
 'DYS_unit':['DYS_unit','dys_unit'],
 'CTS':['CTS','cts','关键期敏感性'],
 'OA':['OA','oa','重合系数'],
 'theta_WS':['theta_WS','θWS','WS权重'],
 'theta_DR':['theta_DR','θDR','DR权重'],
 'theta_SV':['theta_SV','θSV','SV权重'],
 'is_simulation':['is_simulation','模拟数据','simulation'],
}

UNKNOWN_NAMES={'unknown','未知','来源未知','未知来源','未识别','unmapped','unknown source'}

def _is_missing(v):
    if v is None: return True
    if isinstance(v,float) and pd.isna(v): return True
    return str(v).strip()==''

def _to_num(v):
    if _is_missing(v): return None
    try: return float(str(v).replace(',','').strip())
    except Exception: return None

def _is_unknown_row(row) -> bool:
    if bool(row.get('unknown_flag')): return True
    nm=str(row.get('node_name') or '').strip().lower()
    nid=str(row.get('node_id') or '').strip().lower()
    return nm in UNKNOWN_NAMES or nid in {'unknown','unmapped'}

def _quantity_to_tons(value: float, unit: str) -> float:
    u=(unit or '').lower()
    if '万吨' in u: return value*10000
    if u in {'kg','千克','公斤'}: return value/1000
    return value

_CN_DIGITS={'零':0,'一':1,'二':2,'两':2,'三':3,'四':4,'五':5,'六':6,'七':7,'八':8,'九':9,'十':10}
def _cn_simple_number(s: str):
    s=(s or '').strip()
    if s in _CN_DIGITS: return float(_CN_DIGITS[s])
    if len(s)==2 and s[0]=='十' and s[1] in _CN_DIGITS: return 10+_CN_DIGITS[s[1]]
    if len(s)==2 and s[1]=='十' and s[0] in _CN_DIGITS: return _CN_DIGITS[s[0]]*10
    if len(s)==3 and s[1]=='十' and s[0] in _CN_DIGITS and s[2] in _CN_DIGITS:
        return _CN_DIGITS[s[0]]*10+_CN_DIGITS[s[2]]
    return None

def _percent_from_fragment(fragment: str):
    m=re.search(r'(\d+(?:\.\d+)?)\s*%',fragment)
    if m: return float(m.group(1))/100.0
    m=re.search(r'百分之([零一二两三四五六七八九十]{1,3})',fragment)
    if m:
        x=_cn_simple_number(m.group(1))
        return None if x is None else x/100.0
    m=re.search(r'([零一二两三四五六七八九十])成',fragment)
    if m:
        x=_cn_simple_number(m.group(1))
        return None if x is None else x/10.0
    return None

def _param_from_text(text: str, field: str):
    m=re.search(rf'(?<![A-Za-z]){field}\s*[=＝:：]\s*(-?\d+(?:\.\d+)?)',text,re.I)
    return float(m.group(1)) if m else None

def extract_structured_message(text: str) -> pd.DataFrame:
    """Deterministic text normalizer used before AI fallback.

    It preserves 0, explicit Unknown share, quantities, custom parameters and simulation identity.
    """
    msg=(text or '').strip()
    if not msg: return pd.DataFrame(columns=CANONICAL)
    norm=msg.replace('—','-').replace('–','-').replace('－','-')
    mats=[m for m in ['甘蔗','甜菜','大豆'] if m in norm]
    material=mats[0] if len(mats)==1 else ''
    year_m=re.search(r'\b(20\d{2})\s*年?',norm); year=int(year_m.group(1)) if year_m else ''
    ent=''
    em=re.search(r'(?:企业|公司)[：:\s]*([A-Za-z0-9\u4e00-\u9fff（）()\-]{2,30})',norm)
    if em: ent=em.group(1).strip('，,。；; ')

    total_qty=None
    tm=re.search(r'(?:共|合计|总计|总量)\s*(\d+(?:\.\d+)?)\s*(万吨|吨|kg|千克|公斤)',norm,re.I)
    if tm: total_qty=_quantity_to_tons(float(tm.group(1)),tm.group(2))

    # Custom deterministic parameters; 0 is preserved.
    params={f:_param_from_text(norm,f) for f in ['WS','SV','DR','BWD','DYS','CTS','OA']}
    theta={}
    for f in ['WS','DR','SV']:
        pats=[rf'(?:theta[_\s]*{f}|θ\s*{f}|{f}\s*权重)\s*[=＝:：]\s*(\d+(?:\.\d+)?)']
        val=None
        for pat in pats:
            mm=re.search(pat,norm,re.I)
            if mm: val=float(mm.group(1)); break
        theta[f]=val
    scale01=bool(re.search(r'(?:WS|SV).{0,30}(?:0\s*[-~到—]\s*1|0\s*至\s*1).*尺度|(?:0\s*[-~到—]\s*1).*尺度',norm,re.I))
    scale05=bool(re.search(r'(?:WS|SV).{0,30}(?:0\s*[-~到—]\s*5|0\s*至\s*5).*尺度|(?:0\s*[-~到—]\s*5).*尺度',norm,re.I))
    simulation=bool(re.search(r'模拟|测试数据|假设数据|演示数据|仅用于测试',norm))
    confidence='Low' if simulation else 'High'
    source_type='S-user-simulation' if simulation else 'V-user-confirmed'

    # Known node aliases.
    k=known_nodes_df()
    found=[]
    for _,kr in k.iterrows():
        if material and kr['material']!=material: continue
        full=str(kr['node_name'])
        base=full.split('(')[0]
        aliases=[full,base]
        if '-' in base:
            parts=[x for x in base.split('-') if x]
            if parts: aliases.append(parts[0])
            if len(parts)>=2: aliases.append('-'.join(parts[:2]))
        # project-specific high-value aliases used in tests/demos
        if str(kr['node_id'])=='SC01': aliases += ['广西崇左','广西崇左-江州-北海','崇左']
        if str(kr['node_id'])=='SC03': aliases += ['巴西中南部','圣保罗','米纳斯吉拉斯']
        if str(kr['node_id'])=='SC06': aliases += ['澳大利亚昆士兰','昆士兰北部','昆士兰']
        pos=-1; alias=''
        for a in sorted(set(aliases),key=len,reverse=True):
            aa=a.replace('—','-').replace('–','-')
            if aa and aa in norm:
                pos=norm.find(aa); alias=aa; break
        if pos<0: continue
        frag=norm[pos:pos+90]
        w=_percent_from_fragment(frag)
        q=None; qunit=''
        qm=re.search(r'(\d+(?:\.\d+)?)\s*(万吨|吨|kg|千克|公斤)',frag,re.I)
        if qm:
            q=_quantity_to_tons(float(qm.group(1)),qm.group(2)); qunit='吨'
            if w is None and total_qty and total_qty>0: w=q/total_qty
        found.append({
            'enterprise':ent,'material':material or kr['material'],'node_id':kr['node_id'],'node_name':full,
            'purchase_weight':w,'year':year,'purchase_quantity':q,'purchase_unit':qunit,
            'procurement_basis':'数量' if q is not None else ('比例' if w is not None else ''),
            'source_type':source_type,'confidence':confidence,'data_origin':'user_text',
            'is_simulation':simulation,'human_confirmed':False,'unknown_flag':False,
            'evidence':frag[:220],
        })

    # Node-code expressions even without names.
    for mt in re.finditer(r'\b(SC\d{2}|SB\d{2}|SY\d{2})\b',norm,re.I):
        nid=mt.group(1).upper()
        if any(str(x['node_id'])==nid for x in found): continue
        rows=k[k.node_id==nid]
        if rows.empty: continue
        kr=rows.iloc[0]; frag=norm[mt.start():mt.start()+90]
        w=_percent_from_fragment(frag)
        qm=re.search(r'(\d+(?:\.\d+)?)\s*(万吨|吨|kg|千克|公斤)',frag,re.I)
        q=_quantity_to_tons(float(qm.group(1)),qm.group(2)) if qm else None
        if w is None and q is not None and total_qty: w=q/total_qty
        found.append({'enterprise':ent,'material':material or kr.material,'node_id':nid,'node_name':kr.node_name,
                      'purchase_weight':w,'year':year,'purchase_quantity':q,'purchase_unit':'吨' if q is not None else '',
                      'procurement_basis':'数量' if q is not None else ('比例' if w is not None else ''),
                      'source_type':source_type,'confidence':confidence,'data_origin':'user_text','is_simulation':simulation,
                      'human_confirmed':False,'unknown_flag':False,'evidence':frag[:220]})

    # Explicit unknown/unmapped share.
    um=re.search(r'(来源未知|未知来源|未识别(?:来源)?|Unknown)[^0-9零一二三四五六七八九十]{0,12}((?:\d+(?:\.\d+)?)\s*%|百分之[零一二两三四五六七八九十]{1,3}|[一二两三四五六七八九十]成)',norm,re.I)
    if um:
        uw=_percent_from_fragment(um.group(2))
        found.append({'enterprise':ent,'material':material,'node_id':'','node_name':'Unknown','purchase_weight':uw,'year':year,
                      'purchase_quantity':None,'purchase_unit':'','procurement_basis':'比例',
                      'source_type':source_type,'confidence':confidence,'data_origin':'user_text','is_simulation':simulation,
                      'human_confirmed':False,'unknown_flag':True,'evidence':um.group(0)})

    # A fully custom single-node test is allowed when all risk parameters are supplied.
    if not found and any(v is not None for v in params.values()):
        w=None
        wm=re.search(r'(?:采购占比|采购比例|权重)[：:=＝\s]*(\d+(?:\.\d+)?)\s*%',norm)
        if wm: w=float(wm.group(1))/100
        elif re.search(r'100\s*%',norm): w=1.0
        node_name='用户自定义节点'
        found.append({'enterprise':ent,'material':material,'node_id':'USER01','node_name':node_name,'purchase_weight':w,'year':year,
                      'purchase_quantity':None,'purchase_unit':'','procurement_basis':'比例',
                      'source_type':source_type,'confidence':confidence,'data_origin':'user_text','is_simulation':simulation,
                      'human_confirmed':False,'unknown_flag':False,'evidence':norm[:220]})

    # If total quantity was given but individual quantities were parsed, compute missing weights.
    if total_qty and total_qty>0:
        for row in found:
            if not row.get('unknown_flag') and row.get('purchase_weight') is None and row.get('purchase_quantity') is not None:
                row['purchase_weight']=float(row['purchase_quantity'])/total_qty

    # Attach custom parameters to known/custom node rows.
    for row in found:
        if row.get('unknown_flag'): continue
        for f,v in params.items(): row[f]=v
        if params['WS'] is not None:
            row['WS_unit']='normalized_0_1' if scale01 or params['WS']==0 else ('aqueduct_score_0_5' if scale05 or params['WS']>1 else '')
        if params['SV'] is not None:
            row['SV_unit']='normalized_0_1' if scale01 or params['SV']==0 else ('aqueduct_score_0_5' if scale05 or params['SV']>1 else '')
        row['DYS_unit']='score_0_1'
        row['theta_WS']=theta['WS']; row['theta_DR']=theta['DR']; row['theta_SV']=theta['SV']

    if not found:
        return pd.DataFrame(columns=CANONICAL)
    df=pd.DataFrame(found)
    for c in CANONICAL:
        if c not in df.columns: df[c]=None
    return df[CANONICAL]

def known_nodes_df():
    rows=[]
    for nid,h in b.HAZARD_DATA.items():
        mat='甘蔗' if nid.startswith('SC') else ('甜菜' if nid.startswith('SB') else '大豆')
        rows.append({'node_id':nid,'node_name':h['name'],'material':mat,'data_type':h['data_type'],'hazard_confidence':h['confidence'],'pfaf_id':h['pfaf_id']})
    return pd.DataFrame(rows)


def load_demo_procurement(material='甘蔗') -> pd.DataFrame:
    rows=[]
    for nid,e in b.EXPOSURE_DATA[material].items():
        if e['W']<=0: continue
        rows.append({'enterprise':b.ENTERPRISE_MAP[material],'material':material,'node_id':nid,
                     'node_name':b.HAZARD_DATA[nid]['name'],'purchase_weight':e['W'],'year':2025,
                     'source_type':'A','confidence':'Low'})
    return pd.DataFrame(rows)


def _extract_docx(path: str) -> str:
    d=Document(path)
    parts=[p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            parts.append(' | '.join(c.text for c in row.cells))
    return '\n'.join(parts)


def _extract_pdf(path: str) -> str:
    reader=PdfReader(path)
    return '\n'.join((p.extract_text() or '') for p in reader.pages)


def extract_candidates_from_text(text: str) -> pd.DataFrame:
    rows=[]
    mats=[m for m in ['甘蔗','甜菜','大豆'] if m in text]
    default_mat=mats[0] if len(mats)==1 else ''
    for nid,h in b.HAZARD_DATA.items():
        names=[h['name'], h['name'].split('(')[0], h['name'].split('-')[0]]
        pos=-1; matched=''
        for n in names:
            if n and n in text:
                pos=text.find(n); matched=n; break
        if pos<0: continue
        context=text[max(0,pos-80):pos+len(matched)+100]
        pct=re.search(r'(\d+(?:\.\d+)?)\s*%', context)
        w=float(pct.group(1))/100 if pct else None
        mat='甘蔗' if nid.startswith('SC') else ('甜菜' if nid.startswith('SB') else '大豆')
        rows.append({'enterprise':'','material':default_mat or mat,'node_id':nid,'node_name':h['name'],
                     'purchase_weight':w,'year':'','source_type':'V-candidate','confidence':'Low',
                     'evidence':context.replace('\n',' ')[:180]})
    return pd.DataFrame(rows)


def read_user_file(path: str):
    ext=Path(path).suffix.lower()
    if ext=='.csv':
        for enc in ['utf-8-sig','utf-8','gb18030']:
            try: return {'kind':'table','data':pd.read_csv(path,encoding=enc),'text':''}
            except Exception: pass
        raise ValueError('CSV 编码无法识别')
    if ext in ['.xlsx','.xls']:
        return {'kind':'table','data':pd.read_excel(path),'text':''}
    if ext=='.docx':
        text=_extract_docx(path); return {'kind':'text','data':extract_candidates_from_text(text),'text':text}
    if ext=='.pdf':
        text=_extract_pdf(path); return {'kind':'text','data':extract_candidates_from_text(text),'text':text}
    raise ValueError('仅支持 PDF、DOCX、XLSX、CSV')


def suggest_mapping(columns: List[str]) -> pd.DataFrame:
    cols=list(map(str,columns)); rows=[]
    targets=['enterprise','material','node_id','node_name','purchase_weight','purchase_quantity','purchase_unit',
             'procurement_basis','year','WS','WS_unit','SV','SV_unit','DR','BWD','DYS','DYS_unit','CTS','OA',
             'theta_WS','theta_DR','theta_SV','is_simulation']
    lower={c:c.strip().lower() for c in cols}
    for target in targets:
        best=''
        for c in cols:
            if lower[c] in [a.lower() for a in ALIASES.get(target,[target])]:
                best=c; break
        if not best:
            for c in cols:
                if any(a.lower() in lower[c] or lower[c] in a.lower() for a in ALIASES.get(target,[target])):
                    best=c; break
        rows.append({'system_field':target,'source_column':best,
                     'required':target in ['material']})
    return pd.DataFrame(rows)

def _best_node_match(name: str, material: str='') -> Tuple[str,str,float]:
    if not name: return '','',0
    k=known_nodes_df()
    if material in ['甘蔗','甜菜','大豆']:
        k=k[k.material==material]
    name=str(name).strip()
    # exact id/name first
    exact=k[(k.node_id.astype(str)==name)|(k.node_name.astype(str)==name)]
    if not exact.empty:
        r=exact.iloc[0]; return r.node_id,r.node_name,1.0
    scores=[]
    for _,r in k.iterrows():
        base=r.node_name.split('(')[0]
        score=max(difflib.SequenceMatcher(None,name,r.node_name).ratio(),difflib.SequenceMatcher(None,name,base).ratio())
        if name in r.node_name or base in name: score=max(score,0.9)
        scores.append((score,r.node_id,r.node_name))
    score,nid,nm=max(scores,default=(0,'',''))
    return (nid,nm,score) if score>=0.55 else ('','',score)


def apply_mapping(raw: pd.DataFrame, mapping: pd.DataFrame, weight_mode='0-1') -> pd.DataFrame:
    if raw is None or len(raw)==0: return pd.DataFrame(columns=CANONICAL)
    mapping=dict(zip(mapping['system_field'],mapping['source_column'])) if mapping is not None and len(mapping) else {}
    out=pd.DataFrame(index=raw.index)
    for target in [x for x in CANONICAL if x not in {'source_type','confidence','data_origin','human_confirmed','unknown_flag','evidence'}]:
        src=mapping.get(target,'')
        out[target]=raw[src] if src and src in raw.columns else None
    for c in ['purchase_weight','purchase_quantity','WS','SV','DR','BWD','DYS','CTS','OA','theta_WS','theta_DR','theta_SV']:
        out[c]=pd.to_numeric(out[c],errors='coerce')
    if weight_mode=='百分数(0-100)': out['purchase_weight']=out['purchase_weight']/100.0

    # Quantity-only standard inputs are converted deterministically within material/year groups.
    if out['purchase_weight'].isna().all() and out['purchase_quantity'].notna().any():
        keys=['material']
        if out['year'].notna().any(): keys.append('year')
        for _,idx in out.groupby(keys,dropna=False).groups.items():
            total=float(pd.to_numeric(out.loc[idx,'purchase_quantity'],errors='coerce').fillna(0).sum())
            if total>0: out.loc[idx,'purchase_weight']=pd.to_numeric(out.loc[idx,'purchase_quantity'],errors='coerce')/total

    out['source_type']='V-user-confirmed'
    out['confidence']='High'
    out['data_origin']='user_file'
    out['human_confirmed']=True
    out['is_simulation']=out['is_simulation'].fillna(False).astype(bool)
    out['unknown_flag']=False
    for i,row in out.iterrows():
        nm=str(row.get('node_name') or '').strip()
        nid=str(row.get('node_id') or '').strip()
        if nm.lower() in UNKNOWN_NAMES or nid.lower() in {'unknown','unmapped'}:
            out.at[i,'unknown_flag']=True
            out.at[i,'node_id']=''; out.at[i,'node_name']='Unknown'
            continue
        mat=str(row.get('material') or '').strip()
        custom_haz=all(not _is_missing(row.get(f)) for f in ['WS','SV','DR'])
        if nid not in b.HAZARD_DATA and not custom_haz:
            mid,mname,score=_best_node_match(nm or nid,mat)
            if mid:
                out.at[i,'node_id']=mid;out.at[i,'node_name']=mname
                if score<0.9: out.at[i,'confidence']='Medium'
        if custom_haz and not str(out.at[i,'node_id'] or '').strip():
            out.at[i,'node_id']=f'USER{i+1:02d}'
            out.at[i,'node_name']=nm or '用户自定义节点'
    out['evidence']=''
    for c in CANONICAL:
        if c not in out.columns: out[c]=None
    return out[CANONICAL]

def validate_normalized(df: pd.DataFrame) -> dict:
    issues=[]; warnings=[]; gaps=[]
    if df is None or len(df)==0:
        return {'status':'insufficient','issues':['没有可计算的采购记录'],'warnings':[],'data_gaps':['采购数据为空'],
                'coverage':{},'unknown_share_by_material':{}}
    x=df.copy()
    for c in CANONICAL:
        if c not in x.columns: x[c]=None
    if x['material'].replace('',pd.NA).isna().any(): issues.append('存在缺失 material')
    if pd.to_numeric(x['purchase_weight'],errors='coerce').isna().any():
        issues.append('存在缺失/非数值 purchase_weight')
    weights=pd.to_numeric(x['purchase_weight'],errors='coerce')
    if ((weights<0)|(weights>1)).fillna(False).any():
        issues.append('采购份额必须在 0–1 之间')

    invalid_nodes=[]
    for _,r in x.iterrows():
        if _is_unknown_row(r): continue
        nid=str(r.get('node_id') or '').strip()
        custom_haz=all(not _is_missing(r.get(f)) for f in ['WS','SV','DR'])
        if nid not in b.HAZARD_DATA and not custom_haz:
            invalid_nodes.append(str(r.get('node_name') or nid or '未命名节点'))
    if invalid_nodes:
        gaps.append('以下节点无法匹配到当前16节点水风险库，且未提供完整自定义危险度参数：'+', '.join(invalid_nodes))

    # WS/SV scale rule: zero is unambiguous; positive <=1 without unit is ambiguous.
    for idx,r in x.iterrows():
        for f in ['WS','SV']:
            val=_to_num(r.get(f)); unit=str(r.get(f+'_unit') or '').strip()
            if val is None: continue
            if val<0 or val>5: issues.append(f'第{idx+1}行 {f}={val} 超出允许范围')
            if 0<val<=1 and not unit:
                issues.append(f'第{idx+1}行 {f}={val} 未说明是 0–1 还是 0–5 尺度，请确认')
            if unit=='normalized_0_1' and val>1: issues.append(f'第{idx+1}行 {f} 已声明 0–1 尺度但数值>1')
        for f in ['DR','BWD','DYS','CTS','OA']:
            val=_to_num(r.get(f))
            if val is not None and not (0<=val<=1): issues.append(f'第{idx+1}行 {f}={val} 必须在 0–1')

    coverage_by_mat={}; unknown_by_mat={}
    for mat,g in x.groupby('material',dropna=False):
        known=0.0; explicit_unknown=0.0
        for _,r in g.iterrows():
            w=_to_num(r.get('purchase_weight')) or 0.0
            if _is_unknown_row(r): explicit_unknown+=w
            else: known+=w
        total=known+explicit_unknown
        coverage_by_mat[str(mat)]=round(known,6)
        if total>1.0001:
            issues.append(f'{mat} 采购权重合计={total:.4f}（{total*100:.1f}%）>100%，请修改')
        residual=max(0.0,1.0-total)
        unknown=explicit_unknown+residual
        unknown_by_mat[str(mat)]=round(unknown,6)
        if unknown>0:
            warnings.append(f'{mat} 已知可定位采购份额={known:.4f}，Unknown={unknown:.4f} 将保留，不自动归一化')
    status='conflict' if issues else ('partial' if invalid_nodes or warnings else 'success')
    return {'status':status,'issues':list(dict.fromkeys(issues)),'warnings':list(dict.fromkeys(warnings)),
            'data_gaps':list(dict.fromkeys(gaps)),'coverage':coverage_by_mat,
            'unknown_share_by_material':unknown_by_mat}

