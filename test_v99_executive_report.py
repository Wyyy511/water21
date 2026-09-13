from pathlib import Path
import pandas as pd
from tools.data_tool import load_demo_procurement
from tools.baseline_tool import calculate_baseline
from tools.scenario_tool import run_scenario
from tools.report_facts import build_report_facts, fallback_structured_report
from tools.export_tool import export_risk_report
from docx import Document

def test_enterprise_report_hides_backend_jargon():
    df=load_demo_procurement('大豆')
    b=calculate_baseline(df,run_id='v99_report')
    s=run_scenario(b,'NodeFailure','大豆',failure_fraction=1.0,inventory=0.10)
    facts=build_report_facts(baseline=b,scenario=s,validation={},records=df.to_dict(orient='records'),snapshot_id='snap_v99')
    rep=fallback_structured_report(facts)
    path=export_risk_report(b,s,{},'snap_v99',report_facts=facts,structured_report=rep)
    doc=Document(path)
    text='\n'.join(p.text for p in doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            text+='\n'+' | '.join(c.text for c in row.cells)
    forbidden=['PRWI','KnownRisk','Coverage','scored_coverage','NodeFailure','gross_loss','unmet_demand','conditional_PRWI','fact_id','fig1_','fig2_','run_','snapshot','BWD','DYS','CTS','theta_']
    assert not any(x in text for x in forbidden),[x for x in forbidden if x in text]
    for required in ['管理摘要','供应地区风险概览','主要水风险来源','压力测试','AI 风险判断与管理建议','数据可信度与使用说明']:
        assert required in text
