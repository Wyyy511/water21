from __future__ import annotations
from datetime import datetime
from pathlib import Path
import math
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from core.paths import OUTPUT_DIR
from tools.chart_tool import generate_report_charts
from tools.report_facts import fallback_structured_report

CJK="Noto Sans CJK SC"
NAVY="17365D"; BLUE="2F75B5"; LIGHT="EAF3F8"; PALE="F5F9FC"; GREEN="1F8A70"; AMBER="C98A16"; GREY="6B7785"

def _m(v): return v is None or (isinstance(v,float) and (math.isnan(v) or math.isinf(v)))
def _num(v,n=3):
    if _m(v): return "不可计算"
    try:return f"{float(v):.{n}f}"
    except:return "不可计算"
def _pct(v):
    if _m(v):return "不可计算"
    try:return f"{float(v)*100:.1f}%"
    except:return "不可计算"
def _shade(cell,fill=LIGHT):
    pr=cell._tc.get_or_add_tcPr();x=OxmlElement("w:shd");x.set(qn("w:fill"),fill);pr.append(x)
def _margins(cell,top=80,start=100,bottom=80,end=100):
    tc=cell._tc;tcPr=tc.get_or_add_tcPr();mar=tcPr.first_child_found_in("w:tcMar")
    if mar is None: mar=OxmlElement("w:tcMar");tcPr.append(mar)
    for tag,val in [("top",top),("start",start),("bottom",bottom),("end",end)]:
        el=mar.find(qn("w:"+tag))
        if el is None: el=OxmlElement("w:"+tag);mar.append(el)
        el.set(qn("w:w"),str(val));el.set(qn("w:type"),"dxa")
def _cell(cell,text,bold=False,color=None,size=9.5,align=None):
    cell.text="";p=cell.paragraphs[0]
    if align is not None:p.alignment=align
    r=p.add_run(str(text));r.bold=bold;r.font.name=CJK;r.font.size=Pt(size)
    r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),CJK)
    if color:r.font.color.rgb=RGBColor.from_string(color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;_margins(cell)
def _heading(doc,text,level=1):
    p=doc.add_paragraph();r=p.add_run(text);r.bold=True;r.font.name=CJK;r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),CJK)
    r.font.size=Pt(16 if level==1 else 12);r.font.color.rgb=RGBColor.from_string(NAVY);return p
def _para(doc,text,bold_prefix=None):
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(6);p.paragraph_format.line_spacing=1.25
    if bold_prefix and text.startswith(bold_prefix):
        r=p.add_run(bold_prefix);r.bold=True;r.font.color.rgb=RGBColor.from_string(NAVY);p.add_run(text[len(bold_prefix):])
    else:p.add_run(text)
    return p
def _bullet(doc,text):
    p=doc.add_paragraph(style="List Bullet");p.paragraph_format.space_after=Pt(4);p.add_run(text);return p

def _fonts(doc):
    for st in doc.styles:
        try:st.font.name=CJK;st._element.rPr.rFonts.set(qn("w:eastAsia"),CJK)
        except:pass
    for p in doc.paragraphs:
        for r in p.runs:r.font.name=CJK;r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),CJK)

def _rank_label(rank,total):
    if rank==1:return "优先关注"
    if rank<=max(2,(total+1)//2):return "重点监控"
    return "常规监控"

def _risk_driver_summary(summary):
    vals=[("长期水资源紧张",float(summary.get("path_contrib_WS") or 0)),("历史干旱暴露",float(summary.get("path_contrib_DR") or 0)),("季节性供水波动",float(summary.get("path_contrib_SV") or 0))]
    total=sum(v for _,v in vals)
    vals.sort(key=lambda x:x[1],reverse=True)
    return vals,total

def _scenario_plain(sc):
    if not sc or not sc.get("data"):return None
    d=sc.get("data") or {};typ=d.get("scenario_type")
    if typ=="NodeFailure":
        return {"name":"主要供应地区中断","target":d.get("target_node_name") or d.get("target_node_id") or "主要供应地区",
                "shock":_pct(d.get("failure_fraction_f")),"affected":_pct(d.get("gross_loss")),"inventory":_pct(d.get("inventory_used")),
                "replacement":_pct(d.get("replacement_allocated")),"gap":_pct(d.get("unmet_demand"))}
    b=d.get("PRWI_baseline");s=d.get("PRWI_scenario") if d.get("PRWI_scenario") is not None else d.get("PRWI_future")
    return {"name":"水压力变化测试","before":_num(b),"after":_num(s),"change":_num(d.get("PRWI_delta"))}

def _source_names(facts):
    names=[]
    for d in facts.get("sources") or []:
        name=d.get("source_org") or d.get("secondary_source_org")
        if name and name not in names:names.append(str(name))
    return names[:5]

def export_risk_report(baseline_result:dict, scenario_result:dict|None, validation:dict|None, snapshot_id:str="",
                       baseline_explanation:str="", scenario_explanation:str="",
                       report_facts:dict|None=None, structured_report:dict|None=None)->str:
    facts=report_facts or {};report=structured_report or fallback_structured_report(facts)
    charts=generate_report_charts(facts,OUTPUT_DIR,prefix=(snapshot_id or baseline_result.get("run_id") or "report"))
    scope=(facts.get("scope") or [{}])[0];ent=scope.get("enterprise") or "本次分析企业";mat=scope.get("material") or "原材料";year=scope.get("year") or ""
    summaries=(facts.get("baseline") or {}).get("summary") or [];summary=summaries[0] if summaries else {}
    nodes=sorted([x for x in (facts.get("nodes") or []) if x.get("C") is not None],key=lambda x:float(x.get("C") or 0),reverse=True)
    driver_vals,driver_total=_risk_driver_summary(summary) if summary else ([],0)
    top=nodes[0] if nodes else None;sc=_scenario_plain(facts.get("scenarios"))
    quality=facts.get("quality") or {};identity=(quality.get("data_identity") or {}).get("label") or ""
    demo=bool(quality.get("proxy_fields") or quality.get("assumptions") or "代理" in identity or "假设" in identity or "示范" in identity)

    doc=Document();sec=doc.sections[0];sec.top_margin=Inches(.58);sec.bottom_margin=Inches(.58);sec.left_margin=Inches(.68);sec.right_margin=Inches(.68)
    doc.styles["Normal"].font.name=CJK;doc.styles["Normal"].font.size=Pt(10.5)

    # Cover / simple identity
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(26)
    r=p.add_run("WaterPulse");r.bold=True;r.font.size=Pt(13);r.font.color.rgb=RGBColor.from_string(BLUE)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(f"{mat}上游供应链水风险决策报告");r.bold=True;r.font.size=Pt(23);r.font.color.rgb=RGBColor.from_string(NAVY)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"{ent}"+(f" · {year}" if year else "")).italic=True
    if demo:
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        rr=p.add_run("当前为参考 / 示范数据分析，用于风险筛查与方案比较");rr.font.color.rgb=RGBColor.from_string(AMBER);rr.bold=True

    # 1 executive summary
    _heading(doc,"1. 管理摘要")
    _para(doc,str(report.get("executive_summary") or "本次已完成上游供应链水风险分析，以下展示关键结果与管理建议。"))
    cards=doc.add_table(rows=2,cols=2);cards.alignment=WD_TABLE_ALIGNMENT.CENTER
    metric1=f"{_num(summary.get('PRWI'))}\n综合相对水风险指数" if summary else "不可计算\n综合相对水风险指数"
    metric2=(f"{top.get('node_name')}\n贡献约 {_pct(top.get('contribution_share'))}" if top else "不可计算\n首要关注供应地区")
    metric3=(f"{driver_vals[0][0]}\n约占 {(driver_vals[0][1]/driver_total*100):.1f}%" if driver_vals and driver_total else "不可计算\n首要风险因素")
    metric4=(f"剩余缺口 {sc.get('gap')}\n主要供应地区中断测试" if sc and sc.get('gap') else "未运行\n压力测试")
    for cell,text,color in [(cards.cell(0,0),metric1,BLUE),(cards.cell(0,1),metric2,NAVY),(cards.cell(1,0),metric3,GREEN),(cards.cell(1,1),metric4,AMBER)]:
        _shade(cell,PALE);_cell(cell,text,True,color,11,WD_ALIGN_PARAGRAPH.CENTER)

    # 2 supply risk overview
    _heading(doc,"2. 供应地区风险概览")
    _para(doc,"本节只展示企业真正需要用于采购判断的结果：采购集中在哪里、哪些供应地区对整体风险影响最大，以及需要优先管理的顺序。")
    if charts.get("fig1_procurement") and Path(charts["fig1_procurement"]).exists():
        doc.add_picture(charts["fig1_procurement"],width=Inches(6.25));p=doc.add_paragraph("当前采购结构");p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    if nodes:
        t=doc.add_table(rows=1,cols=5);t.style="Table Grid";t.alignment=WD_TABLE_ALIGNMENT.CENTER
        for i,h in enumerate(["供应地区","采购占比","地区相对水风险","对整体风险贡献","管理优先级"]):_cell(t.rows[0].cells[i],h,True,"FFFFFF");_shade(t.rows[0].cells[i],BLUE)
        for rank,n in enumerate(nodes,1):
            vals=[n.get("node_name") or "—",_pct(n.get("weight")),_num(n.get("R"),3),_pct(n.get("contribution_share")),_rank_label(rank,len(nodes))]
            cells=t.add_row().cells
            for i,v in enumerate(vals):_cell(cells[i],v,False,NAVY if i==4 else None)
    if report.get("node_findings"):
        for x in report.get("node_findings")[:3]:
            if x.get("text"):_bullet(doc,str(x.get("text")))

    # 3 drivers
    _heading(doc,"3. 主要水风险来源")
    _para(doc,"风险来源用于回答“为什么需要关注这些供应地区”。这里不展示计算公式，只保留企业需要理解的风险构成。")
    if charts.get("fig2_paths") and Path(charts["fig2_paths"]).exists():
        doc.add_picture(charts["fig2_paths"],width=Inches(6.25));p=doc.add_paragraph("主要水风险来源及其相对贡献");p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    if driver_vals and driver_total:
        for name,val in driver_vals:
            _bullet(doc,f"{name}：约占整体风险的 {val/driver_total*100:.1f}%。")
    for x in report.get("baseline_findings") or []:
        if x.get("text"):_para(doc,str(x.get("text")))

    # 4 scenario
    _heading(doc,"4. 压力测试：如果关键供应地区发生中断")
    if sc:
        if sc.get("gap"):
            _para(doc,f"本次测试假设 {sc.get('target')} 的供应发生 {sc.get('shock')} 中断。直接受影响的采购量约为 {sc.get('affected')}；现有库存可缓冲约 {sc.get('inventory')}，替代供应可覆盖约 {sc.get('replacement')}，最终仍有约 {sc.get('gap')} 的需求无法满足。")
            _para(doc,"这项测试的作用不是预测事件一定会发生，而是帮助企业提前判断：如果主供地区突然失效，当前库存和替代采购能力是否足够。")
        else:
            _para(doc,f"压力测试下，综合相对水风险指数由 {sc.get('before')} 变化为 {sc.get('after')}（变化 {sc.get('change')}）。")
        if charts.get("fig4_scenario") and Path(charts["fig4_scenario"]).exists():
            doc.add_picture(charts["fig4_scenario"],width=Inches(6.1));p=doc.add_paragraph("关键供应地区中断后的供应影响");p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for x in report.get("scenario_findings") or []:
            if x.get("text"):_para(doc,str(x.get("text")))
    else:
        _para(doc,"本次没有运行压力测试。若企业需要评估关键供应地区中断后的库存和替代供应能力，可在 WaterPulse 中继续运行该测试。")

    # 5 AI decision recommendations
    _heading(doc,"5. AI 风险判断与管理建议")
    _para(doc,"以下建议由 AI 基于本次确定性风险结果、供应地区贡献和压力测试结果生成，用于帮助采购、供应链与 ESG 团队确定行动顺序。")
    acts=report.get("actions") or []
    if acts:
        t=doc.add_table(rows=1,cols=5);t.style="Table Grid";t.alignment=WD_TABLE_ALIGNMENT.CENTER
        for i,h in enumerate(["优先级","管理对象","建议动作","为什么优先","建议责任/时点"]):_cell(t.rows[0].cells[i],h,True,"FFFFFF");_shade(t.rows[0].cells[i],NAVY)
        for a in acts:
            owner=" / ".join(x for x in [str(a.get("department") or ""),str(a.get("timing") or "")] if x)
            vals=[a.get("priority") or "—",a.get("object") or "—",a.get("action") or "—",a.get("priority_reason") or "—",owner or "—"]
            cells=t.add_row().cells
            for i,v in enumerate(vals):_cell(cells[i],v,False,NAVY if i==0 else None)
    else:
        _bullet(doc,"优先核实高贡献供应地区的采购占比、供水稳定性与替代来源。")
        _bullet(doc,"将关键地区纳入常态化水风险监测，并把库存和第二供应来源纳入采购预案。")
        _bullet(doc,"在正式调整采购结构前，用企业真实供应商和采购数据替换当前参考或示范参数。")

    # 6 simple quality / use boundary
    _heading(doc,"6. 数据可信度与使用说明")
    if demo:_bullet(doc,"本次分析包含参考、代理或示范数据。结果适合用于风险筛查和方案比较，不应直接作为最终采购决策。")
    cov=summary.get("Coverage") if summary else None;unk=summary.get("unknown_share_U") if summary else None
    if cov is not None:_bullet(doc,f"当前已识别采购来源覆盖约 {_pct(cov)}；未识别采购份额约 {_pct(unk)}。")
    for lim in report.get("limitations") or []:_bullet(doc,str(lim))
    names=_source_names(facts)
    if names:_bullet(doc,"主要参考数据来源包括："+"、".join(names)+"。详细数据版本与追溯信息已保留在系统后台，不在企业版报告中展开。")
    _bullet(doc,"本报告用于支持采购与供应链风险管理；重大决策仍需结合成本、质量、合同、供应商尽调与企业内部判断。")

    _fonts(doc)
    out=OUTPUT_DIR/f"WaterPulse_Enterprise_Decision_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    doc.save(out);return str(out)
